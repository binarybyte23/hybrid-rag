"""Hand-rolled paragraph-aware chunker + bulk insert."""
from __future__ import annotations
import hashlib
import os
from hybrid_rag.embeddings import EmbeddingProvider


def chunk_text(text: str, source: str, max_chars: int = 1000,
               page: int | None = None) -> list[tuple[str, str, str, int | None]]:
    """Split on blank lines, pack paragraphs to ~max_chars. Returns (id, content, source, page)."""
    paras = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 2 <= max_chars:
            cur = (cur + "\n\n" + p).strip()
        else:
            if cur:
                chunks.append(cur)
            cur = p if len(p) <= max_chars else p[:max_chars]
    if cur:
        chunks.append(cur)
    out = []
    for i, c in enumerate(chunks or ([text[:max_chars]] if text.strip() else [])):
        did = hashlib.md5(f"{source}:{page}:{i}:{c[:64]}".encode()).hexdigest()[:12]
        out.append((did, c, source, page))
    return out


def _read_segments(fp: str) -> list[tuple[str, int | None]]:
    """Extract (text, page) segments. Page is PDF page / slide number, else None."""
    low = fp.lower()
    if low.endswith((".txt", ".md", ".markdown")):
        with open(fp, errors="ignore") as f:
            return [(f.read(), None)]
    if low.endswith(".pdf"):
        from pypdf import PdfReader  # lazy: only needed for PDFs
        return [((p.extract_text() or ""), i + 1)
                for i, p in enumerate(PdfReader(fp).pages)]
    if low.endswith(".pptx"):
        from pptx import Presentation  # lazy
        prs = Presentation(fp)
        segs = []
        for i, s in enumerate(prs.slides):
            parts = [sh.text for sh in s.shapes
                     if sh.has_text_frame and sh.text.strip()]
            if s.has_notes_slide:
                parts.append(s.notes_slide.notes_text_frame.text)
            segs.append(("\n\n".join(parts), i + 1))
        return segs
    if low.endswith(".docx"):
        from docx import Document  # lazy
        d = Document(fp)
        return [("\n\n".join([p.text for p in d.paragraphs] +
                             [c.text for t in d.tables for r in t.rows for c in r.cells]), None)]
    return []


def _read_file(fp: str) -> str:
    """Plain-text fallback (tests / unknown callers)."""
    return "\n\n".join(t for t, _ in _read_segments(fp))


async def ingest_paths(paths: list[str], embedder: EmbeddingProvider,
                       max_chars: int = 1000) -> int:
    files: list[str] = []
    for p in paths:
        if os.path.isdir(p):
            files += [os.path.join(p, f) for f in os.listdir(p)
                      if f.lower().endswith((".txt", ".md", ".markdown", ".pdf", ".pptx",
                                             ".key", ".pages", ".ppt", ".doc", ".docx"))]
        elif os.path.exists(p):
            files.append(p)
    rows: list[tuple] = []
    skipped: list[str] = []
    for fp in files:
        got = False
        for text, page in _read_segments(fp):
            if not text.strip():
                continue
            got = True
            for did, content, source, pg in chunk_text(text, os.path.basename(fp), max_chars, page):
                rows.append((did, content, os.path.basename(fp), pg))
        if not got:
            skipped.append(os.path.basename(fp))
    if skipped:
        # ponytail: no native .key/.pages/.ppt reader exists; export first, don't parse zip junk
        print(f"skip: {', '.join(skipped)} — unsupported or empty "
              f"(export Keynote→.pptx, Pages→.pdf/.txt first)")
    if not rows:
        return 0
    vecs = await embedder.embed_texts([r[1] for r in rows])
    from hybrid_rag.db import get_pool  # lazy: no asyncpg needed for chunk_text/tests
    pool = await get_pool()
    async with pool.acquire() as conn:
        await conn.executemany(
            """INSERT INTO chunks (id, content, source, page, embedding)
               VALUES ($1, $2, $3, $4, $5) ON CONFLICT (id) DO UPDATE
               SET content=EXCLUDED.content, page=EXCLUDED.page, embedding=EXCLUDED.embedding""",
            [(r[0], r[1], r[2], r[3], v) for r, v in zip(rows, vecs)])
    return len(rows)
